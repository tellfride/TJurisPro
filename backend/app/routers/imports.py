from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, Response, UploadFile, File
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from ..auth import require_consultant_permission, require_roles
from ..database import get_db
from ..models import Client, Company, Loan, User, UserRole
from ..services.cpf_validator import validate_cpf as _validate_cpf
from ..services.audit_logger import log_action
from ..services.interest_engine import calculate_total_amount, generate_installments, next_loan_number
from ..services.telegram import notify_company

router = APIRouter(prefix="/api/imports", tags=["imports"])

HEADER_FONT = Font(bold=True)

COL_NAME = "Nome do Cliente"
COL_DOCUMENT = "CPF do Solicitante"
COL_PHONE = "Telefone"
COL_EMAIL = "Email"
COL_CEP = "CEP"
COL_ADDRESS = "Endereço"
COL_ADDRESS_NUMBER = "Número"
COL_REF1_NAME = "Referência 1 - Nome"
COL_REF1_PHONE = "Referência 1 - Telefone"
COL_REF2_NAME = "Referência 2 - Nome"
COL_REF2_PHONE = "Referência 2 - Telefone"
COL_REF3_NAME = "Referência 3 - Nome"
COL_REF3_PHONE = "Referência 3 - Telefone"
COL_PRINCIPAL = "Valor Solicitado"
COL_RATE = "Taxa de Juros %"
COL_TERM = "Prazo (meses)"
COL_LATE_FEE = "Multa por Atraso R$/dia"
COL_START_DATE = "Data do Empréstimo (AAAA-MM-DD)"
COL_NOTES = "Observações"

TEMPLATE_HEADERS = [
    COL_NAME, COL_DOCUMENT, COL_PHONE, COL_EMAIL, COL_CEP, COL_ADDRESS, COL_ADDRESS_NUMBER,
    COL_REF1_NAME, COL_REF1_PHONE, COL_REF2_NAME, COL_REF2_PHONE, COL_REF3_NAME, COL_REF3_PHONE,
    COL_PRINCIPAL, COL_RATE, COL_TERM, COL_LATE_FEE, COL_START_DATE, COL_NOTES,
]
REQUIRED_HEADERS = [COL_NAME, COL_DOCUMENT]
LOAN_HEADERS = [COL_PRINCIPAL, COL_RATE, COL_TERM, COL_LATE_FEE]


@router.get("/clients-loans/template.xlsx")
def download_template(
    user: User = Depends(require_roles(UserRole.gestor, UserRole.consultor)),
):
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes e Empréstimos"
    ws.append(TEMPLATE_HEADERS)
    for cell in ws[1]:
        cell.font = HEADER_FONT
    ws.append([
        "Maria Silva", "123.456.789-09", "(11) 99999-0000", "maria@exemplo.com",
        "01310-100", "Avenida Paulista", "1000",
        "José (irmão)", "(11) 98888-0001", "Ana (vizinha)", "(11) 98888-0002", "", "",
        1000, 30, 10, 10, "2026-01-01", "Cliente de exemplo",
    ])
    for column_cells in ws.columns:
        length = max(len(str(c.value)) if c.value is not None else 0 for c in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 12), 40)

    instructions = wb.create_sheet("Instruções")
    instructions.append(["Como preencher esta planilha"])
    instructions["A1"].font = Font(bold=True, size=13)
    lines = [
        "",
        "Uma linha = um cliente (+ opcionalmente um empréstimo lançado para ele).",
        "Colunas obrigatórias para o cliente: " + f"{COL_NAME}, {COL_DOCUMENT}" + ".",
        "Para lançar um empréstimo junto, preencha TODAS as colunas de empréstimo (" +
        f"{COL_PRINCIPAL}, {COL_RATE}, {COL_TERM}, {COL_LATE_FEE}" +
        "). Para importar só o cliente (sem empréstimo), deixe essas 4 colunas em branco.",
        "Se o nome do cliente já existir cadastrado na sua empresa, o cliente é reaproveitado",
        "(não duplica) e só um novo empréstimo é lançado para ele — nesse caso o CPF da linha é ignorado.",
        "CPF deve ter 11 dígitos (com ou sem pontuação). As 3 referências são opcionais.",
        "Endereço e Número são opcionais. No site, o Endereço é preenchido automaticamente a partir do",
        "CEP — na planilha, preencha os dois manualmente (Endereço = rua/bairro/cidade, Número = do imóvel).",
        "Taxa de Juros % é o total do período (ex: 30 = 30% sobre o valor solicitado no prazo todo).",
        "Data do Empréstimo no formato AAAA-MM-DD (ex: 2026-01-31). Se deixar em branco, usa a data de hoje.",
        "Não apague a linha de cabeçalho (linha 1).",
    ]
    for line in lines:
        instructions.append([line])

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="jurispro_modelo_importacao.xlsx"'},
    )


def _clean(value):
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _parse_decimal(value, field_name: str, row_number: int) -> Decimal:
    if value is None or value == "":
        raise ValueError(f"Linha {row_number}: '{field_name}' é obrigatório")
    try:
        return Decimal(str(value).replace(",", "."))
    except (InvalidOperation, ValueError):
        raise ValueError(f"Linha {row_number}: '{field_name}' inválido ({value!r})")


def _parse_int(value, field_name: str, row_number: int) -> int:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        raise ValueError(f"Linha {row_number}: '{field_name}' inválido ({value!r})")


def _parse_start_date(value, row_number: int) -> date:
    if value in (None, ""):
        return date.today()
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip())
    except ValueError:
        raise ValueError(f"Linha {row_number}: 'Data do Empréstimo' inválida ({value!r}), use AAAA-MM-DD")


@router.post("/clients-loans")
async def import_clients_loans(
    file: UploadFile = File(...),
    user: User = Depends(require_roles(UserRole.gestor, UserRole.consultor)),
    db: Session = Depends(get_db),
):
    require_consultant_permission(db, user, "register_clients")
    require_consultant_permission(db, user, "register_loans")
    if not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=422, detail="Envie um arquivo .xlsx")

    content = await file.read()
    try:
        wb = load_workbook(filename=BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(status_code=422, detail="Não foi possível ler o arquivo. Verifique se é um .xlsx válido")

    ws = wb.worksheets[0]
    header_row = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    header_index = {name: idx for idx, name in enumerate(header_row)}
    missing = [h for h in REQUIRED_HEADERS if h not in header_index]
    if missing:
        raise HTTPException(
            status_code=422,
            detail=f"Cabeçalho da planilha inválido. Colunas faltando: {', '.join(missing)}. Baixe o modelo atualizado.",
        )

    def cell(row, header_name):
        idx = header_index.get(header_name)
        return row[idx] if idx is not None and idx < len(row) else None

    errors: list[dict] = []
    clients_created = 0
    loans_created = 0
    next_number = next_loan_number(db, user.company_id)

    company = db.get(Company, user.company_id)
    max_clients = company.max_clients if company else None
    client_count = db.query(Client).filter(Client.company_id == user.company_id).count()

    existing_clients = {
        c.name.strip().lower(): c
        for c in db.query(Client).filter(Client.company_id == user.company_id).all()
    }

    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(v in (None, "") for v in row):
            continue
        savepoint = db.begin_nested()
        try:
            name = _clean(cell(row, COL_NAME))
            if not name:
                raise ValueError(f"Linha {row_number}: '{COL_NAME}' é obrigatório")

            loan_values = {h: cell(row, h) for h in LOAN_HEADERS}
            loan_filled = [h for h, v in loan_values.items() if v not in (None, "")]
            wants_loan = len(loan_filled) > 0
            if wants_loan and len(loan_filled) < len(LOAN_HEADERS):
                missing = [h for h in LOAN_HEADERS if h not in loan_filled]
                raise ValueError(
                    f"Linha {row_number}: preencha todas as colunas de empréstimo ({', '.join(missing)}) "
                    "ou deixe as 4 em branco para importar só o cliente"
                )

            if wants_loan:
                principal = _parse_decimal(loan_values[COL_PRINCIPAL], COL_PRINCIPAL, row_number)
                rate = _parse_decimal(loan_values[COL_RATE], COL_RATE, row_number)
                term_months = _parse_int(loan_values[COL_TERM], COL_TERM, row_number)
                late_fee = _parse_decimal(loan_values[COL_LATE_FEE], COL_LATE_FEE, row_number)
                start_date = _parse_start_date(cell(row, COL_START_DATE), row_number)
                if principal <= 0:
                    raise ValueError(f"Linha {row_number}: '{COL_PRINCIPAL}' deve ser maior que zero")
                if term_months <= 0:
                    raise ValueError(f"Linha {row_number}: '{COL_TERM}' deve ser maior que zero")
                if rate < 0:
                    raise ValueError(f"Linha {row_number}: '{COL_RATE}' não pode ser negativa")
                if late_fee < 0:
                    raise ValueError(f"Linha {row_number}: '{COL_LATE_FEE}' não pode ser negativa")

            client = existing_clients.get(name.lower())
            if client is None:
                if max_clients is not None and client_count >= max_clients:
                    raise ValueError(
                        f"Linha {row_number}: limite de {max_clients} clientes da empresa atingido — "
                        "cliente não importado"
                    )
                raw_document = _clean(cell(row, COL_DOCUMENT))
                if not raw_document:
                    raise ValueError(f"Linha {row_number}: '{COL_DOCUMENT}' é obrigatório para cliente novo")
                try:
                    document = _validate_cpf(raw_document)
                except ValueError as exc:
                    raise ValueError(f"Linha {row_number}: '{COL_DOCUMENT}' {exc} ({raw_document!r})")
                client = Client(
                    company_id=user.company_id,
                    name=name,
                    document=document,
                    phone=_clean(cell(row, COL_PHONE)),
                    email=_clean(cell(row, COL_EMAIL)),
                    cep=_clean(cell(row, COL_CEP)),
                    address=_clean(cell(row, COL_ADDRESS)),
                    address_number=_clean(cell(row, COL_ADDRESS_NUMBER)),
                    notes=_clean(cell(row, COL_NOTES)),
                    reference1_name=_clean(cell(row, COL_REF1_NAME)),
                    reference1_phone=_clean(cell(row, COL_REF1_PHONE)),
                    reference2_name=_clean(cell(row, COL_REF2_NAME)),
                    reference2_phone=_clean(cell(row, COL_REF2_PHONE)),
                    reference3_name=_clean(cell(row, COL_REF3_NAME)),
                    reference3_phone=_clean(cell(row, COL_REF3_PHONE)),
                    created_by=user.id,
                )
                db.add(client)
                db.flush()
                existing_clients[name.lower()] = client
                client_count += 1
                clients_created += 1

            if wants_loan:
                total_amount = calculate_total_amount(principal, rate)
                loan = Loan(
                    company_id=user.company_id,
                    loan_number=next_number,
                    client_id=client.id,
                    principal=principal,
                    interest_rate=rate,
                    term_months=term_months,
                    start_date=start_date,
                    late_fee_per_day=late_fee,
                    total_amount=total_amount,
                    created_by=user.id,
                )
                db.add(loan)
                db.flush()
                for installment in generate_installments(loan):
                    installment.loan_id = loan.id
                    db.add(installment)
                loans_created += 1
                next_number += 1
            savepoint.commit()
        except ValueError as exc:
            savepoint.rollback()
            errors.append({"row": row_number, "message": str(exc)})
        except Exception as exc:  # noqa: BLE001 - report and keep processing remaining rows
            savepoint.rollback()
            errors.append({"row": row_number, "message": f"Erro inesperado: {exc}"})

    log_action(
        db, user, "importar_planilha", "import_batch", None,
        {"clientes_criados": clients_created, "emprestimos_criados": loans_created, "linhas_com_erro": len(errors)},
        company_id=user.company_id,
    )
    db.commit()

    if loans_created > 0:
        notify_company(
            db, user.company_id,
            (
                f"📥 <b>Importação em massa</b>\n"
                f"{clients_created} cliente(s) e {loans_created} empréstimo(s) importados por {user.name}"
                + (f"\n⚠️ {len(errors)} linha(s) com erro" if errors else "")
            ),
        )

    return {"clients_created": clients_created, "loans_created": loans_created, "errors": errors}
