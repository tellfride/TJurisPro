from datetime import date, datetime
from io import BytesIO

from fastapi import APIRouter, Depends, Query, Response
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session, joinedload

from ..auth import require_consultant_permission, require_roles
from ..database import get_db
from ..models import Client, Company, InstallmentStatus, Loan, Payment, User, UserRole
from ..services.analytics import profit_split

router = APIRouter(prefix="/api/reports", tags=["reports"])

# Texto que o Excel/LibreOffice trata como fórmula ou comando quando é o começo da célula.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _append_row_safe(ws, row: list) -> None:
    """ws.append() que nunca deixa texto digitado por usuário virar fórmula.

    O openpyxl grava como FÓRMULA qualquer texto que comece com "=". Como nomes de
    cliente/empresa e observações vêm de digitação (ou de planilha importada), um
    cliente chamado `=HYPERLINK("http://golpe","Clique")` viraria um link/fórmula
    executado na máquina de quem abre o relatório. Aqui a célula é forçada a texto
    e marcada com quotePrefix (o Excel mostra o conteúdo literalmente)."""
    ws.append(row)
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str) and cell.value.startswith(_FORMULA_PREFIXES):
            cell.data_type = "s"
            cell.quotePrefix = True

HEADER_FONT = Font(bold=True)
COLUMNS = [
    "Data do Pagamento", "Empresa", "Cliente", "Empréstimo Nº", "Parcela #", "Tipo",
    "Valor Pago (R$)", "Multa Incluída (R$)", "Lucro de Juros (R$)", "Registrado por", "Observações",
]


@router.get("/transactions.xlsx")
def export_transactions(
    company_id: int | None = Query(default=None),
    date_from: date | None = Query(default=None),
    date_to: date | None = Query(default=None),
    user: User = Depends(require_roles(UserRole.administrador, UserRole.gestor, UserRole.consultor)),
    db: Session = Depends(get_db),
):
    require_consultant_permission(db, user, "view_reports")
    scope_company_id = user.company_id if user.role != UserRole.administrador else company_id

    query = (
        db.query(Payment)
        .join(Loan)
        .options(
            joinedload(Payment.loan).joinedload(Loan.client),
            joinedload(Payment.loan).joinedload(Loan.company),
            joinedload(Payment.installment),
            joinedload(Payment.registered_by_user),
        )
    )
    if scope_company_id is not None:
        query = query.filter(Loan.company_id == scope_company_id)
    if date_from is not None:
        query = query.filter(Payment.payment_date >= date_from)
    if date_to is not None:
        query = query.filter(Payment.payment_date <= date_to)
    payments = query.order_by(Payment.payment_date.asc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Transações"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = HEADER_FONT

    for payment in payments:
        loan = payment.loan
        interest_profit, late_fee_profit = profit_split(payment, loan)
        tipo = "Quitação antecipada" if payment.installment_id is None else "Pagamento de parcela"
        registrado_por = payment.registered_by_user.name if payment.registered_by_user else "-"
        _append_row_safe(ws, [
            payment.payment_date.strftime("%d/%m/%Y"),
            loan.company.name,
            loan.client.name,
            loan.loan_number,
            payment.installment.number if payment.installment else "-",
            tipo,
            float(payment.amount),
            float(payment.late_fee_included),
            round(float(interest_profit), 2),
            registrado_por,
            payment.notes or "",
        ])

    for column_cells in ws.columns:
        length = max(len(str(c.value)) if c.value is not None else 0 for c in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(length + 2, 10), 40)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"jurispro_transacoes_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
